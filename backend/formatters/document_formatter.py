from __future__ import annotations

from io import BytesIO


def format_docx(content: str, document_type: str) -> bytes:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("python-docx is required for DOCX export.") from exc

    document = Document()
    document.add_heading(document_type or "Legal Document", level=0)
    for line in content.splitlines():
        if not line.strip():
            document.add_paragraph()
        else:
            document.add_paragraph(line)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def format_pdf(content: str, document_type: str) -> bytes:
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("reportlab is required for PDF export.") from exc

    buffer = BytesIO()
    story = []
    style = getSampleStyleSheet()
    story.append(Paragraph(document_type or "Legal Document", style["Title"]))
    story.append(Spacer(1, 18))

    for paragraph in content.splitlines():
        if paragraph.strip():
            story.append(Paragraph(paragraph, style["BodyText"]))
        else:
            story.append(Spacer(1, 8))

    doc = SimpleDocTemplate(buffer, pagesize=letter, title=document_type or "Legal Document")
    doc.build(story)
    return buffer.getvalue()
